"""Independent python-docx paragraph fixture; no application serializer imports."""
from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

destination = Path(__file__).parent
document = Document()
paragraph = document.add_paragraph()
paragraph.add_run("  A😀 ")
run = paragraph.add_run("marked")
run.bold = True
run.italic = True
run.underline = True
document.add_paragraph("")
document.add_paragraph(" ")
# Keep the independent tool's exact runs and whitespace preservation. The
# application format adds one scene control and omits page-layout metadata.
body = document.element.body
body.remove(body.sectPr)
control = OxmlElement("w:sdt")
properties = OxmlElement("w:sdtPr")
tag = OxmlElement("w:tag")
tag.set(qn("w:val"), "fixture")
properties.append(tag)
content = OxmlElement("w:sdtContent")
for child in list(body):
    content.append(child)
control.extend([properties, content])
body.append(control)
# The default document's mc:Ignorable describes unused Office metadata.
document.element.attrib.clear()
destination.joinpath("python-docx.xml").write_bytes(
    etree.tostring(document.element, encoding="UTF-8", xml_declaration=True)
)
destination.joinpath("python-docx-version.txt").write_text(
    __import__("docx").__version__ + "\n", encoding="utf-8"
)

# Independently construct complete OPC packages and the documented manifest.
# No application writer, ZIP implementation, or output is used.
import hashlib
import json
import zipfile
from xml.sax.saxutils import escape

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
REL = "http://schemas.openxmlformats.org/package/2006/relationships"
OFFICE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/"
MANIFEST = OFFICE + "custom-properties"
CUSTOM = "http://schemas.openxmlformats.org/officeDocument/2006/custom-properties"
VT = "http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes"

def package(name: str, document_xml: str, paragraphs: list[dict]) -> None:
    body = {"type": "doc", "content": paragraphs}
    canonical = json.dumps(body, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    tag = "wr-scene:" + hashlib.sha256(b"book-fixture:scene-fixture:4").hexdigest()[:32]
    manifest = {"version": 2, "book_id": "book-fixture", "item_id": "scene-fixture",
                "doc_rev": 4, "body_hash": hashlib.sha256(canonical.encode()).hexdigest(),
                "tag": tag, "groups": [], "anchors": []}
    types = '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/>'
    for part, kind in [("document", "document.main"), ("settings", "settings"), ("comments", "comments")]:
        types += f'<Override PartName="/word/{part}.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.{kind}+xml"/>'
    types += '<Override PartName="/docProps/custom.xml" ContentType="application/vnd.openxmlformats-officedocument.custom-properties+xml"/></Types>'
    properties = etree.Element(f"{{{CUSTOM}}}Properties", nsmap={None:CUSTOM,"vt":VT})
    for pid,(name_prop,value) in enumerate([("WritingReviewManifestCount","1"),("WritingReviewManifest0000",json.dumps(manifest, ensure_ascii=False, separators=(",", ":")))],2):
        prop=etree.SubElement(properties,f"{{{CUSTOM}}}property",fmtid="{D5CDD505-2E9C-101B-9397-08002B2CF9AE}",pid=str(pid),name=name_prop)
        etree.SubElement(prop,f"{{{VT}}}lpwstr").text=value
    parts = {
        "[Content_Types].xml": types,
        "_rels/.rels": f'<Relationships xmlns="{REL}"><Relationship Id="r1" Type="{OFFICE}officeDocument" Target="word/document.xml"/><Relationship Id="r2" Type="{MANIFEST}" Target="docProps/custom.xml"/></Relationships>',
        "word/document.xml": document_xml.replace('"fixture"', f'"{tag}"'),
        "word/_rels/document.xml.rels": f'<Relationships xmlns="{REL}"><Relationship Id="r3" Type="{OFFICE}settings" Target="settings.xml"/><Relationship Id="r4" Type="{OFFICE}comments" Target="comments.xml"/></Relationships>',
        "word/settings.xml": f'<w:settings xmlns:w="{W}"><w:trackRevisions/></w:settings>',
        "word/comments.xml": f'<w:comments xmlns:w="{W}"/>',
        "docProps/custom.xml": etree.tostring(properties).decode(),
    }
    with zipfile.ZipFile(destination / f"{name}.docx", "w") as archive:
        for part, value in parts.items():
            info = zipfile.ZipInfo(part, (2026, 9, 25, 12, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            archive.writestr(info, value.encode())
    (destination / f"{name}-body.json").write_text(canonical + "\n", encoding="utf-8")

package("python-docx", destination.joinpath("python-docx.xml").read_text(), [
    {"type": "paragraph", "content": [
        {"type": "text", "text": "  A😀 "},
        {"type": "text", "text": "marked", "marks": [{"type": "em"}, {"type": "strong"}, {"type": "underline"}]}]},
    {"type": "paragraph"},
    {"type": "paragraph", "content": [{"type": "text", "text": " "}]}])
package("handwritten", f'<q:document xmlns:q="{W}"><q:body><q:sdt><q:sdtPr><q:tag q:val="fixture"/></q:sdtPr><q:sdtContent><q:p><q:r><q:t>one</q:t></q:r></q:p></q:sdtContent></q:sdt></q:body></q:document>', [{"type": "paragraph", "content": [{"type": "text", "text": "one"}]}])

# Preserve the actual python-docx package, including its default styles/theme
# and document properties, while attaching the same review manifest contract.
import io
saved = io.BytesIO()
document.save(saved)
with zipfile.ZipFile(saved) as archive:
    full_parts = {name: archive.read(name) for name in archive.namelist()}
with zipfile.ZipFile(destination / "python-docx.docx") as archive:
    strict_parts = {name: archive.read(name) for name in archive.namelist()}
full_parts["word/document.xml"] = strict_parts["word/document.xml"]
full_parts["docProps/custom.xml"] = strict_parts["docProps/custom.xml"]
full_parts["word/comments.xml"] = strict_parts["word/comments.xml"]
for part, tag_name, attributes in [
    ("_rels/.rels", f"{{{REL}}}Relationship", {"Id": "review", "Type": MANIFEST, "Target": "docProps/custom.xml"}),
    ("word/_rels/document.xml.rels", f"{{{REL}}}Relationship", {"Id": "reviewComments", "Type": OFFICE + "comments", "Target": "comments.xml"}),
    ("[Content_Types].xml", "{http://schemas.openxmlformats.org/package/2006/content-types}Override", {"PartName": "/word/comments.xml", "ContentType": "application/vnd.openxmlformats-officedocument.wordprocessingml.comments+xml"}),
]:
    root = etree.fromstring(full_parts[part])
    etree.SubElement(root, tag_name, attributes)
    full_parts[part] = etree.tostring(root)
with zipfile.ZipFile(destination / "python-docx-full.docx", "w") as archive:
    for part, value in full_parts.items():
        info = zipfile.ZipInfo(part, (2026, 9, 25, 12, 0, 0))
        info.compress_type = zipfile.ZIP_DEFLATED
        archive.writestr(info, value)
